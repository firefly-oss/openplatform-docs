"""Behavioral regression tests for technical Markdown brand lint."""
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile


SPEC = importlib.util.spec_from_file_location("brand_check", Path(__file__).parents[1] / "brand_check.py")
brand_check = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = brand_check
SPEC.loader.exec_module(brand_check)


class MarkdownBrandCheckTests(unittest.TestCase):
    def lint(self, text, suffix=".md", final=True):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / ("fixture" + suffix)
            path.write_text(text, encoding="utf-8")
            return brand_check.lint_file(path, final=final)

    def rules(self, text, **kwargs):
        return {finding.rule for finding in self.lint(text, **kwargs)}

    def test_banner_attributes_are_not_percentage_claims(self):
        self.assertEqual([], self.lint('<img src="banner.svg" alt="Firefly OpenPlatform" width="100%">'))

    def test_visible_html_and_image_alt_text_are_checked(self):
        for text in ('<span style="width:100%">FireFly</span>', '<img src="banner.svg" alt="FireFly">',
                     '<img data-alt="safe" src="banner.svg" alt="FireFly">'):
            with self.subTest(text=text):
                self.assertIn("brand", self.rules(text))

    def test_closed_technical_fences_are_not_prose(self):
        for marker, language in (("```", "mermaid"), ("~~~~", "java"), ("```", "http")):
            with self.subTest(language=language):
                self.assertEqual([], self.lint(f'{marker}{language}\nA[TODO] → B[FireFly]\n100% guaranteed compliant\n{marker}\n'))

    def test_unclosed_or_mismatched_fences_do_not_hide_following_prose(self):
        for text in ('```java\nFireFly\n', '```java\n~~~\nFireFly\n'):
            with self.subTest(text=text):
                self.assertIn("brand", self.rules(text))

    def test_prose_after_fence_keeps_original_line_number(self):
        findings = self.lint('```java\nString s = "FireFly";\n```\nFireFly\n')
        self.assertEqual([("brand", "fixture.md: line 4")], [(f.rule, f.where) for f in findings])

    def test_inline_code_supports_matching_backtick_runs(self):
        self.assertEqual([], self.lint('Use `values[TODO] → FireFly` and ``a`100% guaranteed compliant``.'))

    def test_unclosed_inline_code_does_not_hide_prose(self):
        self.assertIn("brand", self.rules('Use `an unfinished span. FireFly'))

    def test_valid_link_labels_are_not_placeholders(self):
        text = '[operating guidance](../operations/index.md) and [Source reference](docs/reference.md).'
        self.assertEqual([], self.lint(text))

    def test_link_destinations_with_nested_parentheses_and_titles_are_ignored(self):
        self.assertEqual([], self.lint('[Guide](https://example.test/FireFly(v1) "TODO 100%")'))

    def test_reference_links_and_definitions_preserve_visible_labels(self):
        text = '[operating guidance][guide] and [Source reference][]\n\n[guide]: https://example.test/FireFly\n[Source reference]: ../100%/TODO\n'
        self.assertEqual([], self.lint(text))
        self.assertIn("brand", self.rules('[FireFly][guide]\n\n[guide]: ../reference.md'))

    def test_link_label_violations_remain_visible(self):
        self.assertIn("brand", self.rules('[FireFly](../guide.md)'))
        self.assertIn("brand", self.rules('[FireFly](https://example.test/guide)'))
        self.assertIn("placeholder", self.rules('[TODO](../guide.md)'))
        self.assertIn("invented-number", self.rules('[100% faster](../guide.md)'))

    def test_malformed_links_do_not_hide_claims(self):
        for text in ('[guide](unfinished 100% guaranteed compliant',
                     '[guide](invalid destination 100% guaranteed compliant)',
                     '[guide]: invalid destination 100% guaranteed compliant'):
            with self.subTest(text=text):
                self.assertIn("invented-number", self.rules(text))

    def test_raw_urls_and_autolinks_do_not_read_as_copy(self):
        self.assertEqual([], self.lint('See https://example.test/FireFly/100%/TODO and <https://example.test/FireFly>.'))

    def test_prose_metrics_claims_placeholders_and_glyphs_still_fail(self):
        cases = {'100% guaranteed compliant': {'invented-number', 'compliance-claim'},
                 'TODO TBD Lorem ipsum [Client]': {'placeholder'},
                 'FireFly platform → outcome': {'brand', 'glyph'},
                 '**100%** faster': {'invented-number'}}
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertTrue(expected <= self.rules(text))

    def test_financial_guarantee_operations_are_not_compliance_promises(self):
        for row in ('| `POST` | `/api/v1/guarantees` | Create Guarantee | `201` |',
                    '| GET | /api/v1/guarantee-records/{id} | Get Guarantee | 200 |'):
            with self.subTest(row=row):
                self.assertEqual([], self.lint(row))

    def test_financial_operation_exception_does_not_hide_marketing_claims(self):
        for text in ('| `POST` | `/api/v1/guarantees` | Create Guarantee | Guaranteed compliant |',
                     '| POST | /api/v1/guarantees | Create Guarantee | 100% faster |',
                     'We guarantee compliance.', 'Create Guarantee for every customer.'):
            with self.subTest(text=text):
                self.assertTrue(self.rules(text) & {'compliance-claim', 'invented-number'})

    def test_financial_product_phrases_remain_allowed(self):
        for text in ('Collateral and guarantees', 'Bank guarantee management', 'Guarantee limits', 'Garantías y avales'):
            with self.subTest(text=text):
                self.assertNotIn('compliance-claim', self.rules(text))

    def test_negative_guarantee_caveat_only_exempts_the_negated_verb(self):
        self.assertNotIn('compliance-claim', self.rules('A generated SDK does not guarantee a durable business result.'))
        self.assertIn('compliance-claim', self.rules('The SDK does not guarantee availability. We guarantee compliance.'))

    def test_html_delimiters_in_code_do_not_hide_visible_claims(self):
        for text in ('`<!--` We guarantee compliance. `-->`',
                     '`<code>` We guarantee compliance. `</code>`'):
            with self.subTest(text=text):
                self.assertIn('compliance-claim', self.rules(text))

    def test_malformed_html_block_openers_fail_open(self):
        for tag in ('code', 'pre', 'style', 'script', 'svg'):
            with self.subTest(tag=tag):
                self.assertIn('compliance-claim', self.rules(
                    f'<{tag} title="unfinished>We guarantee compliance.</{tag}>'))

    def test_image_alt_is_not_taken_from_another_attributes_value(self):
        self.assertIn('brand', self.rules("<img title=' alt=\"safe\"' alt=\"FireFly\">"))

    def test_inline_html_preserves_visible_word_adjacency(self):
        self.assertIn('brand', self.rules('Fire<span>Fly</span>'))
        self.assertIn('compliance-claim', self.rules('We guaran<span>tee</span> compliance.'))
        self.assertIn('invented-number', self.rules('100<span>%</span> faster'))

    def test_multiline_image_alt_keeps_its_source_line(self):
        findings = self.lint('<img\n src="image.svg"\n alt="FireFly"\n>\nFireFly')
        self.assertEqual(['fixture.md: line 3', 'fixture.md: line 5'],
                         [f.where for f in findings if f.rule == 'brand'])

    def test_code_delimiters_in_destinations_stay_inside_destinations(self):
        for text in ('[Guide](relative.md "a`") We guarantee compliance. `code`',
                     'https://example.test/a` We guarantee compliance. `code`'):
            with self.subTest(text=text):
                self.assertIn('compliance-claim', self.rules(text))

    def test_multiline_inline_code_is_not_prose(self):
        self.assertEqual([], self.lint('Use `FireFly\n100% compliant` in a shell.'))

    def test_closed_html_metadata_does_not_hide_following_prose(self):
        self.assertEqual([], self.lint('<style>/* TODO */ .x { width: 100%; }</style>\n<!-- TODO -->'))
        self.assertIn('brand', self.rules('<!-- unfinished comment\nFireFly'))


class ExistingFormatTests(unittest.TestCase):
    def test_plaintext_and_json_keep_existing_bracket_and_code_behavior(self):
        for suffix, content in (('.txt', '`FireFly` [Client] 100%'),
                                ('.json', json.dumps({'text': '`FireFly` [Client] 100%'}))):
            with self.subTest(suffix=suffix), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / ('fixture' + suffix)
                path.write_text(content, encoding='utf-8')
                self.assertTrue({'brand', 'placeholder', 'invented-number'} <= {f.rule for f in brand_check.lint_file(path, final=True)})

    def test_html_visible_text_still_checked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'fixture.html'
            path.write_text('<p>FireFly 100% compliant</p>', encoding='utf-8')
            self.assertTrue({'brand', 'invented-number', 'compliance-claim'} <= {f.rule for f in brand_check.lint_file(path)})

    def test_docx_copy_is_not_treated_as_markdown(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'fixture.docx'
            text = '`FireFly` [Client] 100%'
            if importlib.util.find_spec('docx'):
                import docx
                document = docx.Document()
                document.add_paragraph(text)
                document.save(path)
            else:
                with zipfile.ZipFile(path, 'w') as document:
                    document.writestr('word/document.xml', '<w:p><w:t>' + text + '</w:t></w:p>')
            self.assertTrue({'brand', 'placeholder', 'invented-number'} <= {f.rule for f in brand_check.lint_file(path, final=True)})


if __name__ == '__main__':
    unittest.main()
