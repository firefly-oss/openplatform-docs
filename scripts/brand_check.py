#!/usr/bin/env python3
"""Firefly brand and content-policy lint.

Flags risky phrases in slides, documents and copy before anything leaves
the building: invented metrics, client and vendor names, compliance
promises, AI-stance wording that overpromises autonomy, Latin American
Spanish where Spain Spanish is expected, stray language, brand misspellings,
unreplaced placeholders and glyphs Manrope does not have.

Standalone: Python 3.9+ standard library; PyYAML and python-docx are used
when present to read .yaml and .docx files.

  python brand_check.py deck.yaml doc.docx copy.md [--locale es] [--final] [--json]

Exit code 1 when any error is found (warnings alone exit 0 unless --strict).
"""
from __future__ import annotations

import argparse
import html
import json
import re
import sys
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path

# ------------------------------------------------------------------ rules


@dataclass
class Finding:
    rule: str
    severity: str  # error | warning
    message: str
    text: str
    where: str = ""
    suggestion: str = ""


# Lines that cite sources, or are clearly meta, may name publishers and numbers.
SOURCE_KEYS = {"source", "sources", "footnotes", "references", "fuente", "fuentes", "notes"}
SOURCE_PREFIX = re.compile(r"^\s*(\^?\d+\s+)?(source|sources|fuente|fuentes|note|nota|illustrative|ilustrativ)", re.I)

# TODO, TBD and XXX only in capitals: "todo" is an everyday Spanish word.
PLACEHOLDER = re.compile(r"\[[^\]\n]{1,60}\]|Nombre Apellido|\bTBD\b|\bTODO\b|\bXXX\b|(?i:lorem ipsum)")

# Numbers that read as performance claims. Placeholders like [x] or [00]% are fine.
METRIC_PATTERNS = [
    (re.compile(r"(?<![\[\w])\d+(?:[.,]\d+)?\s?%"), "a percentage"),
    (re.compile(r"(?<![\[\w])\d+(?:[.,]\d+)?\s?(?:x|×)\b(?!\s?\d)", re.I), "a multiplier"),
    (re.compile(r"[€$£]\s?\d|\d\s?(?:€|EUR|USD|GBP)\b"), "a price or amount"),
    (re.compile(r"\b\d+(?:[.,]\d+)?\s?(?:million|millions|billion|bn|mn|millones|mil millones|M€|k€)\b", re.I),
     "a money or volume figure"),
    (re.compile(r"\b(?:up to|hasta un|hasta el)\s+\d+", re.I), "an 'up to' claim"),
    (re.compile(r"\b\d+\s?(?:times|veces)\s+(?:faster|cheaper|more|less|más|menos)", re.I), "a comparative claim"),
    (re.compile(r"\b\d+\+?[ \u00a0]+(?:clients|customers|banks|insurers|institutions|countries|projects|"
                r"clientes|bancos|aseguradoras|entidades|países|proyectos)\b", re.I), "a track-record count"),
]
CLAIM_WORDS = re.compile(r"\b(faster|cheaper|reduc(?:e|es|ed|tion)|increase[sd]?|improve[sd]?|saving[s]?|ROI|"
                         r"más rápido|reducción|ahorro|aumento|mejora)\b", re.I)

CLIENTS = ["Santander", "BBVA", "CaixaBank", "Sabadell", "Bankinter", "ING", "Unicaja", "Abanca", "Kutxabank",
           "Ibercaja", "Cajamar", "Mapfre", "Allianz", "AXA", "Zurich", "Generali", "Mutua Madrileña", "Línea Directa",
           "Santalucía", "Catalana Occidente", "Reale", "Barclays", "HSBC", "JPMorgan", "J.P. Morgan", "Deutsche Bank",
           "BNP Paribas", "Société Générale", "UBS", "Lloyds", "NatWest", "Revolut", "N26", "Openbank", "EVO Banco",
           "Aviva", "Prudential", "MetLife", "Sanitas", "Adeslas", "DKV", "Ocaso", "Liberty Seguros", "Plus Ultra"]
VENDORS = ["SAP", "Oracle", "Salesforce", "Temenos", "Finastra", "Mambu", "Thought Machine", "10x Banking",
           "Guidewire", "Duck Creek", "Sapiens", "Majesco", "FIS", "Fiserv", "Jack Henry", "nCino", "Pega",
           "Pegasystems", "Appian", "ServiceNow", "Microsoft", "AWS", "Amazon Web Services", "Azure", "Google Cloud",
           "IBM", "Accenture", "McKinsey", "BCG", "Bain", "Deloitte", "PwC", "EY", "KPMG", "Capgemini", "Infosys",
           "TCS", "OpenAI", "Anthropic", "Claude", "ChatGPT", "Gemini", "Snowflake", "Databricks", "Experian",
           "Equifax", "DocuSign", "Stripe", "Adyen", "Visa", "Mastercard", "Kafka", "Confluent", "Camunda"]

COMPLIANCE = [
    (r"\b(fully |100\s?% )?compliant\b", "Compliance remains the institution's responsibility; describe controls and evidence instead."),
    (r"\bguarantee[sd]?\b|\bguaranteed\b", "Avoid guarantees; describe what the engagement delivers and how it is tested."),
    (r"\bcertified\b|\bcertification\b", "Do not claim certifications."),
    (r"\b(ISO\s?27001|SOC\s?2|PCI[\s-]?DSS|ISO\s?9001|ENS)\b", "Do not cite certifications the company has not published."),
    (r"\bensures? (?:regulatory )?compliance\b", "Compliance remains the institution's responsibility."),
    (r"\b(approved|endorsed) by (the )?(ECB|EBA|EIOPA|Bank of Spain|Banco de España|DGSFP|CNMV|regulator)", "Never imply regulatory approval."),
    (r"\bcumple(n)? (con )?(la )?normativa\b|\bcumplimiento garantizado\b", "El cumplimiento es responsabilidad de cada entidad."),
    (r"\bgarantiza(mos|n)?\b|\bgarantizad[oa]s?\b", "Evita las garantías; describe controles y evidencias."),
    (r"\bcertificad[oa]s?\b|\bhomologad[oa]s?\b", "No afirmes certificaciones."),
]
# Guarantees as a financial instrument (LendingOS "Collateral and guarantees",
# "Garantías y avales") are product vocabulary, not compliance promises.
FINANCIAL_GUARANTEES = re.compile(
    r"\b(?:collateral and guarantees?|guarantees? and collateral|(?:bank|loan|credit|financial|public|state|"
    r"personal|corporate|performance) guarantees?|guarantee (?:schemes?|limits?|lines?|registry|management|"
    r"contracts?|fees?|calls?)|garant[ií]as? y avales|avales y garant[ií]as?)\b", re.I)
COMPLIANCE_OK = re.compile(r"no certification|not .*certif|compliance remains|responsibility of each institution|"
                           r"no compliance guarantee|sin garantías de cumplimiento|responsabilidad de cada entidad|"
                           r"no afirm|no certif", re.I)

AI_STANCE = [
    (r"\bfully autonomous\b|\btotalmente autónom", "AI decides only within approved limits; say so."),
    (r"\bautonomous (credit|underwriting|claims|financial) decisions?\b", "Qualify: within approved limits, explained, recorded and open to human review."),
    (r"\bAI (replaces|will replace|replacing)\b|\breplace (your |the )?(people|staff|underwriters|analysts|teams)\b",
     "AI prepares and decides within limits; it does not replace people with delegated authority."),
    (r"\bla IA (sustituye|reemplaza|sustituirá)\b|\bsustituir a (las personas|los analistas|los suscriptores)",
     "La IA decide dentro de límites aprobados; no sustituye a las personas con facultades."),
    (r"\b(without|no) human (review|oversight|intervention|involvement)\b|\bsin (intervención|revisión|supervisión) humana\b",
     "Every AI decision stays open to human review."),
    (r"\blights[- ]out\b|\bhuman[- ]free\b|\b100\s?% automat|\btotalmente automatizad", "Overpromises automation."),
    (r"\bAI never decides\b|\bla IA nunca decide\b",
     "Outdated stance: AI may decide within approved limits where regulation and risk appetite allow."),
    (r"\bAI decides\b(?![^.]{0,80}(limit|límite|approved|aprobad|appetite|apetito|review|revisión))",
     "State the limits: within approved limits, explained, recorded and open to human review."),
]

LATAM = {
    r"\bcomputadora(s)?\b": "ordenador", r"\bcelular(es)?\b": "móvil", r"\baplicativo(s)?\b": "aplicación",
    r"\bagendar\b": "programar / fijar", r"\breporte(s)?\b": "informe", r"\bchequear\b|\bchecar\b": "comprobar",
    r"\brentar\b": "alquilar", r"\bcosto(s)?\b": "coste", r"\bvideo(s)?\b": "vídeo", r"\blaptop\b": "portátil",
    r"\bfinanciamiento\b": "financiación", r"\baccesar\b": "acceder", r"\bplática\b|\bplaticar\b": "conversación",
    r"\bcarro(s)?\b": "coche", r"\bmanejar\b": "gestionar / conducir", r"\bustedes\b": "vosotros / tú (tuteo, como en la web)",
    r"\be-?mail\b": "correo electrónico", r"\bdeliverables?\b": "entregables", r"\bstakeholders?\b": "partes interesadas",
    r"\bjugar un rol\b": "desempeñar un papel", r"\bal interior de\b": "dentro de", r"\bendeudamiento\b": None,
}
EN_STOP = {"the", "and", "with", "our", "your", "for", "this", "that", "from", "into", "will", "which", "what",
           "how", "are", "is", "of", "to", "we", "it", "be", "on"}
ES_STOP = {"el", "la", "los", "las", "de", "del", "que", "y", "para", "con", "una", "un", "por", "como", "más",
           "sus", "tu", "tus", "nuestro", "nuestra", "se", "es", "en", "al"}
ALLOW_EN_IN_ES = {"front", "middle", "back", "office", "openplatform", "industry", "suite", "firefly", "lendingos",
                  "insuranceos", "riskos", "framework", "foundation", "software", "api", "apis", "attitude", "makes",
                  "difference", "the", "solutions", "inc", "core", "sprint", "workshop", "kick-off", "dashboard"}
SLOGAN = "Attitude makes the difference."
BRAND = [
    (r"\bFireFly\b|\bFIREFLY\b", "Firefly"),
    (r"\bLending OS\b|\bLendingOs\b|\bLending-OS\b", "LendingOS"),
    (r"\bInsurance OS\b|\bInsuranceOs\b", "InsuranceOS"),
    (r"\bRisk OS\b|\bRiskOs\b", "RiskOS"),
    (r"\bOpen Platform\b|\bOpenplatform\b", "OpenPlatform"),
    (r"\bIndustry suite\b", "Industry Suite"),
    (r"\bla actitud marca la diferencia\b|\bla actitud hace la diferencia\b", SLOGAN + " (the slogan stays in English)"),
]
GLYPHS = re.compile(r"[→←↔⇒✓✔✗✘★☆●■◆▶►]")
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿]")


def _word_hits(text, names):
    out = []
    for n in names:
        if re.search(r"(?<![\w-])" + re.escape(n) + r"(?![\w-])", text):
            out.append(n)
    return out


def check_text(text: str, where: str = "", *, locale: str | None = None, final: bool = False,
               is_source: bool = False) -> list[Finding]:
    f: list[Finding] = []
    if not text or not text.strip():
        return f
    t = text
    source_line = is_source or bool(SOURCE_PREFIX.match(t))
    # Placeholders.
    for m in PLACEHOLDER.finditer(t):
        tok = m.group(0)
        if tok.lower() in ("lorem ipsum",) or tok.upper() in ("TODO", "XXX", "TBD"):
            f.append(Finding("placeholder", "error", f"Unfinished text '{tok}'", t, where))
        else:
            f.append(Finding("placeholder", "error" if final else "warning",
                             f"Placeholder '{tok}' must be replaced before the final version", t, where))
    plain = PLACEHOLDER.sub(" ", t)
    # Invented metrics.
    if not source_line:
        for rx, what in METRIC_PATTERNS:
            for m in rx.finditer(plain):
                f.append(Finding("invented-number", "error", f"Looks like {what}: '{m.group(0).strip()}'. Use a "
                                 "placeholder ([x]) unless it comes from a cited source.", t, where,
                                 "Replace with [x] and add a source line, or remove."))
        if CLAIM_WORDS.search(plain) and re.search(r"(?<![\[\w])\d", plain) and not re.search(r"\b(19|20)\d\d\b", plain):
            f.append(Finding("invented-number", "warning", "A number next to a performance word reads as a claim.",
                             t, where))
    # Names.
    if not source_line:
        for n in _word_hits(plain, CLIENTS):
            f.append(Finding("client-name", "error", f"Client or bank name '{n}': never name customers without approval.",
                             t, where, "Use [Client]."))
        for n in _word_hits(plain, VENDORS):
            f.append(Finding("vendor-name", "warning", f"Vendor name '{n}': the website names no vendors; use a "
                             "generic description (core banking, policy administration…).", t, where))
    # Compliance promises.
    if not COMPLIANCE_OK.search(t):
        unguarded = FINANCIAL_GUARANTEES.sub(" ", plain)
        for rx, advice in COMPLIANCE:
            m = re.search(rx, unguarded, re.I)
            if m:
                f.append(Finding("compliance-claim", "error", f"Compliance or certification promise '{m.group(0)}'.",
                                 t, where, advice))
    # AI stance.
    for rx, advice in AI_STANCE:
        m = re.search(rx, plain, re.I)
        if m:
            sev = "warning" if "Outdated" in advice or "State the limits" in advice else "error"
            f.append(Finding("ai-stance", sev, f"AI wording '{m.group(0)}'.", t, where, advice))
    # Brand.
    for rx, fix in BRAND:
        m = re.search(rx, plain, re.I if "actitud" in rx else 0)
        if m:
            f.append(Finding("brand", "error", f"Write '{fix}', not '{m.group(0)}'.", t, where, fix))
    # Spain Spanish.
    if locale == "es":
        for rx, fix in LATAM.items():
            m = re.search(rx, plain, re.I)
            if m and fix:
                f.append(Finding("es-ES", "warning", f"'{m.group(0)}' is not Spain Spanish.", t, where, fix))
    # Stray language.
    words = re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]+", plain.replace(SLOGAN, ""))
    lw = [w.lower() for w in words]
    if locale == "es" and len(lw) >= 4:
        en = [w for w in lw if w in EN_STOP and w not in ALLOW_EN_IN_ES]
        es = [w for w in lw if w in ES_STOP]
        if len(en) >= 2 and len(en) > len(es):
            f.append(Finding("language", "error", "English text in a Spanish deck.", t, where))
    if locale == "en" and len(lw) >= 4:
        es = [w for w in lw if w in ES_STOP - {"a", "un", "es", "se", "en", "al"}]
        en = [w for w in lw if w in EN_STOP]
        if len(es) >= 2 and len(es) > len(en) and "nombre apellido" not in t.lower():
            f.append(Finding("language", "error", "Spanish text in an English deck.", t, where))
    if GLYPHS.search(t):
        f.append(Finding("glyph", "error", f"'{GLYPHS.search(t).group(0)}' is not in Manrope; draw it as a shape or icon.",
                         t, where))
    if EMOJI.search(t):
        f.append(Finding("emoji", "error", "No emoji in brand materials.", t, where))
    if "!" in plain and not plain.strip().startswith("!"):
        f.append(Finding("tone", "warning", "Exclamation mark: the brand voice is calm and precise.", t, where))
    return f


# ------------------------------------------------------------------ readers

def _blank(text):
    """Mask syntax without changing the source line numbers."""
    return re.sub(r"[^\n]", " ", text)


def _balanced_end(text, start, opening, closing):
    depth = 0
    quote = None
    i = start
    while i < len(text):
        char = text[i]
        if char == "\\":
            i += 2
            continue
        if opening == "(" and char in ("\"", "'"):
            if quote == char:
                quote = None
            elif quote is None:
                quote = char
        if quote is None:
            if char == opening:
                depth += 1
            elif char == closing:
                depth -= 1
                if depth == 0:
                    return i + 1
        i += 1
    return None


def _reference_key(label):
    return " ".join(label.split()).casefold()


def _markdown_destination(text):
    return bool(re.fullmatch(
        r'''(?:<[^<>\n]*>|[^\s<>]+)(?:[ \t]+(?:"[^"]*"|'[^']*'|\([^()]*\)))?[ \t]*''', text.strip()))


def _markdown_inline(text, references):
    """Read Markdown syntax in source order so delimiters stay local."""
    tag = re.compile(r'''</?([A-Za-z][\w:-]*)(?:\s+[\w:-]+(?:\s*=\s*(?:"[^"]*"|'[^']*'|[^\s"'=<>`]+))?)*\s*/?>''')
    out = []
    i = 0
    while i < len(text):
        if text[i] == "\\" and i + 1 < len(text):
            if text[i + 1] in r"\`*{}[]()#+-.!_|><":
                out.append(text[i + 1])
            else:
                out.append(text[i:i + 2])
            i += 2
            continue
        if text[i] == "`":
            run = re.match(r"`+", text[i:]).group(0)
            end = re.search(r"(?<!`)" + re.escape(run) + r"(?!`)", text[i + len(run):])
            if end:
                stop = i + len(run) + end.end()
                out.append(_blank(text[i:stop]))
                i = stop
                continue
            out.append(run)
            i += len(run)
            continue
        image = text.startswith("![", i)
        label_start = i + 1 if image else i
        if text[label_start:label_start + 1] == "[":
            label_end = _balanced_end(text, label_start, "[", "]")
            if label_end:
                label = text[label_start + 1:label_end - 1]
                stop = None
                if text[label_end:label_end + 1] == "(":
                    stop = _balanced_end(text, label_end, "(", ")")
                    if stop:
                        destination = text[label_end + 1:stop - 1].strip()
                        if destination and not _markdown_destination(destination):
                            stop = None
                elif text[label_end:label_end + 1] == "[":
                    ref_end = _balanced_end(text, label_end, "[", "]")
                    if ref_end:
                        ref = text[label_end + 1:ref_end - 1] or label
                        if _reference_key(ref) in references:
                            stop = ref_end
                elif _reference_key(label) in references:
                    stop = label_end
                if stop:
                    out.append(_markdown_inline(label, references))
                    out.append(_blank(text[label_end:stop]))
                    i = stop
                    continue
        url = re.match(r"<?(?:https?|ftp)://[^\s<>]+>?", text[i:])
        if url:
            out.append(_blank(url.group(0)))
            i += url.end()
            continue
        if text.startswith("<!--", i):
            end = text.find("-->", i + 4)
            if end != -1:
                stop = end + 3
                out.append("\n" * text[i:stop].count("\n"))
                i = stop
                continue
        match = tag.match(text, i) if text[i] == "<" else None
        if match:
            markup = match.group(0)
            name = match.group(1).lower()
            if name in ("script", "style", "svg", "pre", "code") and not markup.startswith("</"):
                end = re.search(r"</" + re.escape(name) + r"\s*>", text[match.end():], re.I)
                if end:
                    stop = match.end() + end.end()
                    out.append(_blank(text[i:stop]))
                    i = stop
                    continue
            attributes = re.finditer(
                r'''\s+([\w:-]+)(?:\s*=\s*(?:"([^"]*)"|'([^']*)'|([^\s"'=<>`]+)))?''', markup)
            alt = next((attr for attr in attributes if attr.group(1).lower() == "alt"), None)
            if name == "img" and alt and any(value is not None for value in alt.groups()[1:]):
                group = next(j for j in range(2, 5) if alt.group(j) is not None)
                out.append("\n" * markup[:alt.start(group)].count("\n") + alt.group(group)
                           + "\n" * markup[alt.end(group):].count("\n"))
            else:
                # Inline tags join their surrounding text; block tags separate it.
                separator = " " if name in ("br", "hr", "p", "div", "li", "td", "th", "tr", "blockquote",
                                            "h1", "h2", "h3", "h4", "h5", "h6") else ""
                out.append(separator + "\n" * markup.count("\n"))
            i = match.end()
            continue
        out.append(text[i])
        i += 1
    return "".join(out)


def _markdown_api_row(line):
    """Exclude contract identifiers and exact financial operation noun labels."""
    if not line.lstrip().startswith("|"):
        return line
    cells = line.split("|")
    values = [cell.strip().strip("`") for cell in cells]
    if not any(re.fullmatch(r"GET|POST|PUT|PATCH|DELETE|HEAD|OPTIONS", value) for value in values):
        return line
    paths = [i for i, value in enumerate(values)
             if re.fullmatch(r"/[^\s|]*\bguarantee(?:s|-records)(?:/[^\s|]*)?", value)]
    if not paths:
        return line
    for i in paths:
        cells[i] = _blank(cells[i])
    operation = re.compile(r"(?:Create|Get|List|Retrieve|Find|Search|Update|Delete) Guarantees?", re.I)
    for i, value in enumerate(values):
        if operation.fullmatch(value):
            cells[i] = _blank(cells[i])
    return "|".join(cells)


def markdown_items(raw):
    """Yield visible prose from technical Markdown, preserving source lines.

    This is a lint profile, not a renderer. Only complete syntax is masked:
    malformed fences, tags and links fail open rather than hiding later copy.
    """
    lines = raw.splitlines(keepends=True)
    i = 0
    while i < len(lines):
        fence = re.match(r"^ {0,3}(`{3,}|~{3,})([^\n]*)", lines[i])
        if fence and not (fence.group(1)[0] == "`" and "`" in fence.group(2)):
            marker = fence.group(1)
            closer = re.compile(r"^ {0,3}" + re.escape(marker[0]) + "{" + str(len(marker)) + r",}[ \t]*(?:\n)?$")
            end = next((j for j in range(i + 1, len(lines)) if closer.fullmatch(lines[j])), None)
            if end is not None:
                for j in range(i, end + 1):
                    lines[j] = _blank(lines[j])
                i = end + 1
                continue
        i += 1
    raw = "".join(_markdown_api_row(line) for line in lines)
    lines = raw.splitlines()
    references = set()
    for i, line in enumerate(lines):
        ref = re.match(r'^ {0,3}\[([^]\n]+)\]:[ \t]*(.+)$', line)
        if ref and _markdown_destination(ref.group(2)):
            references.add(_reference_key(ref.group(1)))
            lines[i] = _blank(line)
    raw = _markdown_inline("\n".join(lines), references)
    for i, line in enumerate(raw.splitlines(), 1):
        text = html.unescape(line)
        # Emphasis and heading delimiters must not separate a number from '%'.
        text = re.sub(r"[*_]+", "", text)
        text = re.sub(r"^ {0,3}#{1,6}\s+", "", text)
        # Exempt the negated verb, not every other promise on the same line.
        text = re.sub(r"\b(?:does|do|did|will|can) not guarantee\b", " ", text, flags=re.I)
        if text.strip():
            yield f"line {i}", text, False


def strings_in(obj, path=""):
    """(path, text, is_source) for every string in a YAML/JSON structure."""
    if isinstance(obj, str):
        key = path.split(".")[-1].split("[")[0]
        yield path, obj, key in SOURCE_KEYS or any(k in path.split(".") for k in SOURCE_KEYS)
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if k in ("layout", "variant", "theme", "locale", "id", "icon", "pair", "status", "tracker", "density",
                     "accent", "language", "src", "image", "waivers", "waive", "purpose"):
                continue
            yield from strings_in(v, f"{path}.{k}" if path else str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from strings_in(v, f"{path}[{i}]")


def read_file(path: Path):
    """Yield (where, text, is_source) and detect the locale when possible."""
    suffix = path.suffix.lower()
    locale = None
    items = []
    if suffix in (".yaml", ".yml", ".json"):
        raw = path.read_text(encoding="utf-8")
        if suffix == ".json":
            data = json.loads(raw)
        else:
            try:
                import yaml  # type: ignore
                data = yaml.safe_load(raw)
            except ImportError:
                data = {"text": raw}
        if isinstance(data, dict):
            locale = data.get("locale")
            if not locale and isinstance(data.get("meta"), dict):
                locale = data["meta"].get("locale")
        slides = data.get("slides") if isinstance(data, dict) else None
        if isinstance(slides, list):
            for i, s in enumerate(slides, 1):
                for p, t, src in strings_in(s):
                    items.append((f"slide {i} {p}", t, src))
            for p, t, src in strings_in({k: v for k, v in data.items() if k != "slides"}):
                items.append((p, t, src))
        else:
            for p, t, src in strings_in(data):
                items.append((p, t, src))
    elif suffix == ".docx":
        try:
            import docx  # type: ignore
            d = docx.Document(str(path))
            for i, p in enumerate(d.paragraphs):
                items.append((f"paragraph {i + 1}", p.text, p.style.name.lower().startswith("caption")))
            for ti, tbl in enumerate(d.tables):
                for ri, row in enumerate(tbl.rows):
                    for ci, cell in enumerate(row.cells):
                        items.append((f"table {ti + 1} r{ri + 1}c{ci + 1}", cell.text, False))
            for si, sec in enumerate(d.sections):
                for p in sec.header.paragraphs + sec.footer.paragraphs:
                    items.append((f"header/footer {si + 1}", p.text, False))
            lang = re.search(rb'w:lang w:val="([a-z]{2})', zipfile.ZipFile(path).read("word/styles.xml"))
            if lang:
                locale = lang.group(1).decode()
        except ImportError:
            xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
            for i, para in enumerate(re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)):
                items.append((f"paragraph {i + 1}", "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", para)), False))
    elif suffix in (".html", ".htm"):
        raw = path.read_text(encoding="utf-8")
        m = re.search(r'<html[^>]*lang="([a-z]{2})', raw)
        locale = m.group(1) if m else None
        raw = re.sub(r"<(script|style)\b.*?</\1>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<svg\b.*?</svg>", " ", raw, flags=re.S | re.I)
        for i, block in enumerate(re.split(r"<(?:p|h[1-6]|li|td|th|div|figcaption|blockquote)\b[^>]*>", raw)):
            txt = html.unescape(re.sub(r"<[^>]+>", " ", block))
            txt = re.sub(r"\s+", " ", txt).strip()
            if txt:
                items.append((f"block {i}", txt, False))
    elif suffix in (".md", ".markdown"):
        items.extend(markdown_items(path.read_text(encoding="utf-8")))
    else:
        for i, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if line.strip():
                items.append((f"line {i}", line, False))
    return items, locale


def lint_file(path, locale=None, final=False) -> list[Finding]:
    path = Path(path)
    items, detected = read_file(path)
    loc = locale or detected
    out = []
    for where, text, is_source in items:
        for f in check_text(text, f"{path.name}: {where}", locale=loc, final=final, is_source=is_source):
            out.append(f)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("files", nargs="+")
    ap.add_argument("--locale", choices=["en", "es"])
    ap.add_argument("--final", action="store_true", help="placeholders are errors")
    ap.add_argument("--strict", action="store_true", help="warnings fail too")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    findings = []
    for f in a.files:
        findings += lint_file(f, a.locale, a.final)
    if a.json:
        print(json.dumps([asdict(x) for x in findings], ensure_ascii=False, indent=1))
    else:
        for x in findings:
            print(f"{x.severity.upper():7} {x.rule:16} {x.where}: {x.message}" + (f"  → {x.suggestion}" if x.suggestion else ""))
        e = sum(x.severity == "error" for x in findings)
        w = sum(x.severity == "warning" for x in findings)
        print(f"{e} error(s), {w} warning(s) in {len(a.files)} file(s)")
    bad = any(x.severity == "error" for x in findings) or (a.strict and findings)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
